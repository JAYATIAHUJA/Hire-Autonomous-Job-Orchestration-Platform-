"""Resume file sniffing and text extraction (PDF via pypdf, DOCX via python-docx).

The file type comes from the bytes, never from the client's filename or Content-Type.
"""

import io
import zipfile
from dataclasses import dataclass

MAX_RESUME_BYTES = 5 * 1024 * 1024
# A PDF whose text layer yields less than this is treated as scanned/image-only and is
# sent to the model as a document instead.
MIN_TEXT_CHARS = 200
# Far beyond any real resume; protects the prompt budget from padded files.
MAX_TEXT_CHARS = 60_000

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class UnsupportedDocument(ValueError):
    """Not a PDF or DOCX, or not readable."""


@dataclass(frozen=True)
class ExtractedText:
    content_type: str
    text: str  # "" when the text layer is empty

    @property
    def needs_pdf_fallback(self) -> bool:
        return self.content_type == PDF and len(self.text) < MIN_TEXT_CHARS


def sniff_content_type(data: bytes) -> str:
    if data.startswith(b"%PDF-"):
        return PDF
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/document.xml" in archive.namelist():
                    return DOCX
        except zipfile.BadZipFile:
            pass
    raise UnsupportedDocument("Upload a PDF or DOCX resume.")


def _normalise(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise UnsupportedDocument("This PDF is password-protected.")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except PdfReadError as exc:
        raise UnsupportedDocument("This PDF could not be read.") from exc


def _docx_text(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises a variety of zip/xml errors
        raise UnsupportedDocument("This DOCX could not be read.") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(data: bytes) -> ExtractedText:
    content_type = sniff_content_type(data)
    raw = _pdf_text(data) if content_type == PDF else _docx_text(data)
    text = _normalise(raw)
    if len(text) > MAX_TEXT_CHARS:
        raise UnsupportedDocument("This resume has far more text than a resume should; trim it and retry.")
    if content_type == DOCX and not text:
        raise UnsupportedDocument("This DOCX has no text.")
    return ExtractedText(content_type=content_type, text=text)
